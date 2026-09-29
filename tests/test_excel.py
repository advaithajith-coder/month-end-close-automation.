"""The Excel export: its formulas must reproduce the engine's figures.

openpyxl stores formulas without calculating them, so a tiny evaluator (SUM ranges, cell references,
+ and -) recomputes each formula here, the way Excel would.
"""
import re

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402

from close_engine.excel import export_year  # noqa: E402

REF = re.compile(r"\b([A-Z]{1,3})(\d+)\b")
SUM = re.compile(r"SUM\(([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)\)")


def evaluate(ws, ref: str) -> float:
    v = ws[ref].value
    if not (isinstance(v, str) and v.startswith("=")):
        return float(v or 0)
    expr = v[1:]

    def expand(m):
        c1, r1, c2, r2 = column_index_from_string(m[1]), int(m[2]), column_index_from_string(m[3]), int(m[4])
        return "(" + "+".join(f"{get_column_letter(c)}{r}" for c in range(c1, c2 + 1) for r in range(r1, r2 + 1)) + ")"

    expr = SUM.sub(expand, expr)
    expr = REF.sub(lambda m: f"({evaluate(ws, m[0])})", expr)
    assert re.fullmatch(r"[0-9eE.+\-() ]*", expr), expr
    return eval(expr)  # arithmetic only, checked above


@pytest.fixture(scope="module")
def workbook(closed_year, year, tmp_path_factory):
    path = tmp_path_factory.mktemp("xl") / "fy2026.xlsx"
    export_year(closed_year, year, str(path))
    return openpyxl.load_workbook(path)


CENT = 0.005  # dollars are floats in Excel; half a cent absorbs binary rounding noise, never a real difference


def find_row(ws, label: str) -> int:
    return next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=1).value == label)


def test_sheets(workbook):
    assert workbook.sheetnames == ["Summary", "Income Statement", "Balance Sheet", "Trial Balance", "Journal", "Close Log"]


def test_income_statement_formulas_match_the_engine(workbook, year):
    ws = workbook["Income Statement"]
    ni_row, total_col = find_row(ws, "Net income"), get_column_letter(len(year["months"]) + 2)
    assert evaluate(ws, f"{total_col}{ni_row}") == pytest.approx(year["totals"]["netIncome"] / 100, abs=CENT)
    for j, m in enumerate(year["months"], start=2):
        assert evaluate(ws, f"{get_column_letter(j)}{ni_row}") == pytest.approx(m["st"]["netIncome"] / 100, abs=CENT)
    assert evaluate(ws, f"{total_col}{find_row(ws, 'Net sales')}") == pytest.approx(year["totals"]["netSales"] / 100, abs=CENT)
    assert evaluate(ws, f"{total_col}{find_row(ws, 'Gross profit')}") == pytest.approx(year["totals"]["grossProfit"] / 100, abs=CENT)


def test_balance_sheet_balances_in_excel(workbook):
    ws = workbook["Balance Sheet"]
    row = find_row(ws, "Check: assets − (liabilities + equity)")
    assert evaluate(ws, f"B{row}") == pytest.approx(0, abs=CENT)
    assert evaluate(ws, f"C{row}") == pytest.approx(0, abs=CENT)


def test_trial_balance_and_journal_totals(workbook):
    ws = workbook["Trial Balance"]
    diff_row = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Difference (must be 0)")
    assert evaluate(ws, f"C{diff_row}") == pytest.approx(0, abs=CENT)
    ws = workbook["Journal"]
    last = ws.max_row
    assert evaluate(ws, f"H{last}") == pytest.approx(evaluate(ws, f"I{last}"), abs=CENT)


def test_summary_equity_check_is_zero(workbook):
    ws = workbook["Summary"]
    row = find_row(ws, "= Difference (must be 0)")
    assert evaluate(ws, f"B{row}") == pytest.approx(0, abs=CENT)
